# поток: card
"""Aggregate saved live MS data and independent supplier-source inventory. No API writes."""
import csv
import json
import math
import sys
import zipfile
import hashlib
from collections import defaultdict, Counter
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from invoice_bot.supplier_groups import GROUPS
from core.ms_api import meta_id
OUT=ROOT/'docs/reports/card_all_supplier_dimensions_2026-10-15'
DIMS=['a883238f-6012-11ed-0a80-0225000bddc2','00c43da1-4f03-11ee-0a80-082e001b3a61','187a0c24-4f03-11ee-0a80-01f1001b6bd8']

def load(name):return json.loads((OUT/name).read_text())
def save(name,data):(OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2))
def positive(value):
    try:return math.isfinite(float(str(value).replace(',','.'))) and float(str(value).replace(',','.'))>0
    except (ValueError,TypeError):return False
def present(value):return value is not None and str(value).strip()!=''
def csvout(name,rows):
    with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

def main():
    products=load('products.json')
    names={s['id']:s['name'] for s in load('suppliers.json')}
    groups={sid:group for group,items in GROUPS.items() for _,_,sid in items}
    aliases={'Суперфайн':'Рапид / Суперфайн','Спринт':'Солюшнс Принт','Спринт МСК':'Солюшнс Принт'}
    stock={k:{r['assortmentId']:r['stock'] for r in load(f'stock_current_{k}.json')} for k in ['remote','zvezdny','discover']}
    allrows=[]
    for p in products:
        sid=meta_id(p,'supplier') or ''
        supplier=names.get(sid,'Без поставщика' if not sid else sid)
        group=groups.get(sid,supplier)
        group=aliases.get(group,group)
        attrs={a.get('id'):a.get('value') for a in p.get('attributes',[])}
        dims=[attrs.get(k) for k in DIMS]
        qty={k:stock[k].get(p['id'],0) for k in stock}
        full=all(present(v) for v in dims)
        valid=all(positive(v) for v in dims)
        any_stock=any(v>0 for v in qty.values())
        allrows.append({'group':group,'supplier':supplier,'supplier_id':sid,'ms_id':p['id'],
            'article':p.get('article',''),'code':p.get('code',''),'external_code':p.get('externalCode',''),
            'name':p.get('name',''),'length_cm':dims[0],'width_cm':dims[1],'height_cm':dims[2],
            'all_three_filled':full,'all_three_positive':valid,'partial_dimensions':any(present(v) for v in dims) and not full,
            'remote_stock':qty['remote'],'zvezdny_stock':qty['zvezdny'],'discover_stock':qty['discover'],
            'any_target_stock':any_stock,'stock_and_three_filled':any_stock and full})
    def aggregate(key):
        result=[]
        for val in sorted({r[key] for r in allrows}):
            rows=[r for r in allrows if r[key]==val]
            label=(rows[0]['supplier']+' ['+val+']') if key=='supplier_id' and val else val
            if key=='supplier_id' and not val: label='Без поставщика'
            result.append({'supplier':label,'total':len(rows),'remote':sum(r['remote_stock']>0 for r in rows),
                'zvezdny':sum(r['zvezdny_stock']>0 for r in rows),'discover':sum(r['discover_stock']>0 for r in rows),
                'any_stock':sum(r['any_target_stock'] for r in rows),'dimensions':sum(r['all_three_filled'] for r in rows),
                'stock_dimensions':sum(r['stock_and_three_filled'] for r in rows),
                'stock_missing_dimensions':sum(r['any_target_stock'] and not r['all_three_filled'] for r in rows),
                'partial_dimensions':sum(r['partial_dimensions'] for r in rows),
                'filled_but_not_positive_numeric':sum(r['all_three_filled'] and not r['all_three_positive'] for r in rows)})
        return sorted(result,key=lambda r:(-r['total'],r['supplier']))
    grouped=aggregate('group');legal=aggregate('supplier_id')
    totals={k:sum(r[k] for r in grouped) for k in grouped[0] if k!='supplier'}
    ids={p['id'] for p in products}
    stock_without_active={k:len({pid for pid,v in data.items() if v>0}-ids) for k,data in stock.items()}
    assert len(ids)==len(products)==sum(r['total'] for r in grouped)
    assert all(r['stock_dimensions']<=r['any_stock']<=r['total'] and r['dimensions']<=r['total'] for r in grouped)
    summary={'at_utc':datetime.now(timezone.utc).isoformat(),'totals':totals,'groups':grouped,
        'legal_suppliers':len(names),'legal_summary_rows':len(legal),'archived_excluded':load('archived_count.json')['count'],
        'stock_ids_outside_active_products':stock_without_active,
        'product_collection':load('products_collection.json'),'stock_collection':load('stocks_collection.json'),
        'units':'unique active MS product IDs; stock > 0 per warehouse, no reserve subtraction',
        'dimensions_definition':'all three custom fields nonempty; physical correctness not asserted'}
    save('summary.json',summary);csvout('supplier_summary.csv',grouped);csvout('supplier_legal_summary.csv',legal);csvout('products_review.csv',allrows)

    manifest=load('source_manifest.json')
    # Verify Excel originals against the uploaded ZIP, not zero-byte failed extractions.
    with zipfile.ZipFile(ROOT/'incoming/size.zip') as archive:
        members={hashlib.sha256(archive.read(i)).hexdigest():i.filename for i in archive.infolist() if not i.is_dir()}
    for s in manifest:
        if s.get('sha256') in members:
            s['original_archive']='incoming/size.zip';s['original_archive_member']=members[s['sha256']]
            s['byte_identical_to_original_archive']=True
    save('source_manifest.json',manifest)

    wb=openpyxl.Workbook();ws=wb.active;ws.title='По поставщикам'
    titles=['Поставщик','Всего активных карточек','Удаленный >0','Звездный >0','Дисквер >0','Есть на любом из 3','Заполнены ДШВ всего','С остатком и ДШВ','С остатком без полных ДШВ','Частично заполнены ДШВ','ДШВ заполнены, но не 3 положительных числа']
    def table(sheet,headers,rows):
        sheet.append(headers)
        for row in rows:
            # Treat supplier-provided strings as text, never Excel formulas.
            sheet.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in row])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B')
        for i in range(1,len(headers)+1):sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width=36 if i==1 else 22
    table(ws,titles,[list(r.values()) for r in grouped]+[['ИТОГО']+list(totals.values())])
    table(wb.create_sheet('Юрлица МС'),titles,[list(r.values()) for r in legal])
    table(wb.create_sheet('Карточки МС'),list(allrows[0]),[list(r.values()) for r in allrows])
    srcrows=[[s['supplier'],s['file'],s['source_kind'],s['status'],s['article_rows'],s['three_positive_rows'],s['unique_articles_with_three'],s.get('source_date',''),s.get('byte_identical_to_original_archive',False)] for s in manifest]
    table(wb.create_sheet('Исходники'),['Поставщик','Файл','Тип','Статус','Строк с артикулом','Строк с 3 положительными размерами','Уникальных артикулов с 3 размерами','Дата внутри источника','Побайтово совпадает с исходным ZIP'],srcrows)
    methods=wb.create_sheet('Методика')
    for line in [
        'Срез 15.10.2026. Только активные entity/product; архивные, услуги и комплекты как отдельные сущности не включены.',
        'Счётчик = уникальные UUID карточек МС, не количество штук и не уникальные внешние коды ТК.',
        'Поставщик из карточки МС. Юрлица объединены по справочнику supplier_groups; обе группы Спринта объединены.',
        'Положительный физический stock; резерв не вычитается. Объединение 3 складов без двойного счёта.',
        'Полные размеры = непустые Длина, см.; Ширина, см.; Высота, см. Числовая проверка отдельным счётчиком.',
        'Заполненность не доказывает соответствие поставщику или реальному коробу. Исторические расчёты габаритов не использованы.',
        'Источник поставщика = явные исходные поля. Даже 3 положительных числа ещё не гарантируют отсутствие опечатки.',
        'Кактус: только ИУ, не МК. Сакура: коробка, не длина/ширина ленты. У Изипринта на листе Чернила единицы не указаны.',
        'Спринт JSON/JSONL — извлечения страниц с URL, не оригинальные HTML. 14.10 две аномальные строки 5519/5875 (9×88×305 мм) остаются HOLD.',
        'Рапид: метры. SF-T9451Bk/3065 остаётся HOLD; каталог содержит также товары, не сопоставленные с МС.',
        'Профилайн: старый XLS имеет вес/объём, а не ДШВ; новые поля API требуют отдельного разбора единиц/осей, не применяются.',
        'incoming/size/ содержит пустые неудачные распаковки; incoming/size_fixed.zip повреждён. Рабочие Excel в incoming/gab совпадают с исходным incoming/size.zip.',
        'Свежесть архивных Excel не подтверждена; сравнение новых прайсов и подготовка записи не выполнялись.',
    ]:methods.append([line])
    methods.column_dimensions['A'].width=150
    wb.save(OUT/'supplier_dimensions_summary.xlsx')

    bundle=OUT/'supplier_original_dimension_sources.zip'
    selected=[s for s in manifest if s['supplier'] in ['Изипринт','Кактус','Сакура'] or (s['supplier']=='Рапид' and '2026-10-14/rapid.json' in s['file'])]
    with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
        for s in selected:
            path=ROOT/s['file'];z.write(path,arcname=(s['supplier']+'_'+path.name))
        z.writestr('source_manifest.json',json.dumps(selected,ensure_ascii=False,indent=2))
        z.writestr('README.txt','Исходные файлы поставщиков, без изменения ячеек и без расчётных размеров.\n3 архивных Excel + API Рапида 14.10.2026.\nУ Кактуса индивидуальная упаковка (ИУ), не мастер-короб (МК).\nУ Изипринта на листе Чернила единицы размеров не указаны.\nТри положительных размера не являются физической проверкой товара; запись в МС не разрешена этим реестром.\n')
    print(json.dumps({'totals':totals,'groups':len(grouped),'legal_suppliers':len(legal),'stock_outside_active':stock_without_active},ensure_ascii=False))

if __name__=='__main__':main()
