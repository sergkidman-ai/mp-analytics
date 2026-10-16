# поток: card
"""Offline proposal only: direct supplier dimensions, no MS writes or scheduling."""
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from prices import blacklist
AUDIT=ROOT/'docs/reports/card_all_supplier_dimensions_2026-10-15'
OUT=ROOT/'docs/reports/card_dimensions_wave1_plan_2026-10-15'
GROUPS={'Кактус':'Кактус','Сакура':'Сакура','Изипринт':'Изипринт','Рапид / Суперфайн':'Рапид'}

def read_csv(path):
    with path.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def key(v):return str(v or '').strip().upper()
def decimal(v):
    try:
        x=Decimal(str(v).replace(',','.').strip())
        return x if x.is_finite() and x>0 else None
    except InvalidOperation:return None
def strongly_different(a,b):
    # Rotate boxes for comparison only; writing retains the supplier's original axis order.
    return any(abs(x-y)>min(x,y)*Decimal('.15') for x,y in zip(sorted(a),sorted(b)))
def fmt(d):return [format(v.normalize(),'f') for v in d]
def save(name,data):(OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2))
def csvout(name,rows):
    if not rows:return
    with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

def main():
    OUT.mkdir(exist_ok=True)
    cards=read_csv(AUDIT/'products_review.csv')
    sources=read_csv(AUDIT/'source_rows.csv')
    idx=defaultdict(list)
    for s in sources:
        supplier=s['supplier']
        if supplier not in GROUPS.values():continue
        if supplier=='Рапид' and '2026-10-14/rapid.json' not in s['file']:continue
        if s['units'] not in ['мм','см','м'] or s['three_positive_values']!='True':continue
        raw=json.loads(s['raw_dimensions'])
        # Excel numeric values have 15 significant digits; remove serialization tails only.
        if s['file'].endswith('.xlsx'):
            raw=[format(decimal(v),'.15g') if decimal(v) else v for v in raw]
        dims=[decimal(v) for v in raw]
        if len(dims)!=3 or not all(dims):continue
        factor={'мм':Decimal('.1'),'см':Decimal(1),'м':Decimal(100)}[s['units']]
        s=dict(s,cm=[v*factor for v in dims]);idx[(supplier,key(s['article']))].append(s)
    blacklist_set=blacklist.load_set()
    # Stop flags come from the saved current MS snapshot, not inferred from external codes.
    flagged={}
    for p in json.loads((AUDIT/'products.json').read_text()):
        reasons=[]
        if blacklist.norm(p.get('article')) in blacklist_set:reasons.append('BLACKLIST')
        if re.search(r'\bбрак\b',p.get('name',''),re.I):reasons.append('DAMAGED_CARD')
        if blacklist.in_damaged(p.get('name')):reasons.append('DAMAGED_PACKAGE')
        if any('стоп' in a.get('name','').lower() and str(a.get('value')).lower() in ['true','да','1'] for a in p.get('attributes',[])):reasons.append('STOP_LIST')
        if reasons:flagged[p['id']]=reasons
    byext=defaultdict(list)
    for c in cards:
        if c['external_code']:byext[c['external_code']].append(c)
    review=[]; ready=[];conflicts=[]
    for c in cards:
        if c['group'] not in GROUPS:continue
        if c['all_three_filled']=='True':continue
        source=idx.get((GROUPS[c['group']],key(c['article'])),[])
        signatures={tuple(s['cm']) for s in source}
        reasons=list(flagged.get(c['ms_id'],[]))
        if re.search(r'\b(?:набор|комплект)\w*\b',c['name'],re.I):
            reasons.append('SET_PACKAGE_SEPARATE_REVIEW')
        if GROUPS[c['group']]=='Рапид' and not key(c['article']).startswith(('SF','OEM')):
            reasons.append('RAPID_ORIGINALS_PREVIOUSLY_EXCLUDED')
        if not source:reasons.append('NO_EXPLICIT_DIRECT_SOURCE')
        elif len(signatures)!=1:reasons.append('AMBIGUOUS_SOURCE_DIMENSIONS')
        dims=source[0]['cm'] if len(signatures)==1 else None
        peer_source_count=peer_ms_count=0
        source_conflicts=ms_conflicts=0
        if not c['external_code']:reasons.append('EMPTY_EXTERNAL_CODE')
        if dims:
            # Conservative first-wave review boundaries, not a claim that every outlier is wrong.
            if min(dims)<1 or max(dims)>100 or dims[0]*dims[1]*dims[2]>25000:
                reasons.append('FIRST_WAVE_SIZE_REVIEW')
            if GROUPS[c['group']]=='Рапид' and key(c['article'])=='SF-T9451BK':reasons.append('KNOWN_SUSPECT_3065')
            for other in byext[c['external_code']]:
                if other['group']==c['group']:continue
                if other['ms_id'] in flagged:continue
                peer_group=GROUPS.get(other['group'])
                peer_sources=idx.get((peer_group,key(other['article'])),[]) if peer_group else []
                for ps in peer_sources:
                    peer_source_count+=1
                    if strongly_different(dims,ps['cm']):
                        source_conflicts+=1
                        conflicts.append({'ms_id':c['ms_id'],'external_code':c['external_code'],'supplier':c['group'],'article':c['article'],
                            'own_cm':' × '.join(fmt(dims)),'other_supplier':other['group'],'other_article':other['article'],
                            'other_ms_id':other['ms_id'],'other_cm':' × '.join(fmt(ps['cm'])),'evidence':'DIRECT_SUPPLIER_SOURCE',
                            'other_source':ps['file']+' / '+ps['sheet']+' / '+str(ps['row_or_item'])})
                other_dims=[decimal(other[k]) for k in ['length_cm','width_cm','height_cm']]
                if all(other_dims):
                    peer_ms_count+=1
                    if strongly_different(dims,other_dims):
                        ms_conflicts+=1
                        conflicts.append({'ms_id':c['ms_id'],'external_code':c['external_code'],'supplier':c['group'],'article':c['article'],
                            'own_cm':' × '.join(fmt(dims)),'other_supplier':other['group'],'other_article':other['article'],
                            'other_ms_id':other['ms_id'],'other_cm':' × '.join(fmt(other_dims)),'evidence':'MS_EXISTING_NOT_VERIFIED',
                            'other_source':'Existing MS fields only; conflict blocks proposal but is not treated as truth'})
            if source_conflicts:reasons.append('CROSS_SUPPLIER_SOURCE_CONFLICT')
            if ms_conflicts:reasons.append('CROSS_SUPPLIER_MS_CONFLICT')
        status='HOLD' if reasons else 'PROPOSED'
        row={'supplier':c['group'],'ms_id':c['ms_id'],'external_code':c['external_code'],'article':c['article'],'name':c['name'],
            'stock_positive':c['any_target_stock']=='True','dimensions_cm':' × '.join(fmt(dims)) if dims else '',
            'status':status,'reasons':';'.join(reasons),'source_rows':len(source),'peer_source_rows':peer_source_count,
            'peer_ms_cards':peer_ms_count,'source_conflicts':source_conflicts,'ms_conflicts':ms_conflicts,
            'source_file':source[0]['file'] if source else '', 'source_sheet':source[0]['sheet'] if source else '',
            'source_row':source[0]['row_or_item'] if source else '',
            'evidence_without_peer':'SINGLE_SUPPLIER_SOURCE' if dims and not peer_source_count and not peer_ms_count else ''}
        review.append(row)
        if status=='PROPOSED':
            ready.append({'ms_id':c['ms_id'],'supplier':c['group'],'supplier_id':c['supplier_id'],'article':c['article'],
                'external_code':c['external_code'],'stock_positive':row['stock_positive'],'dimensions_cm':fmt(dims),
                'source_file':row['source_file'],'source_sheet':row['source_sheet'],'source_row':row['source_row'],
                'all_matching_source_pointers':[{'file':s['file'],'sheet':s['sheet'],'row':s['row_or_item']} for s in source]})
    summary={}
    for group in GROUPS:
        rows=[r for r in review if r['supplier']==group];planned=[r for r in rows if r['status']=='PROPOSED']
        summary[group]={'empty_cards':len(rows),'with_direct_source':sum(r['source_rows']>0 for r in rows),
            'proposed':len(planned),'proposed_stock':sum(r['stock_positive'] for r in planned),
            'hold':sum(r['status']=='HOLD' for r in rows),'reason_counts':dict(Counter(x for r in rows for x in r['reasons'].split(';') if x)),
            'proposed_without_peers':sum(bool(r['evidence_without_peer']) for r in planned)}
    save('summary.json',{'status':'PROPOSAL_NOT_AUTHORIZED','source_snapshot':'2026-10-15 11:43 MSK',
        'supplier_summary':summary,'cards':len(ready),'fields':len(ready)*3,'stock_cards':sum(r['stock_positive'] for r in ready),
        'comparison_rule':'sorted physical sides; strong if any side difference >15% of smaller side; no absolute tolerance',
        'first_wave_review_boundaries':'hold min<1cm, max>100cm, volume>25L; conservative proposal, not physical diagnosis',
        'scope':'all active cards in 4 selected supplier groups with empty dimensions, independent of stock',
        'writes':0,'scheduled':False})
    save('proposed_cards.json',ready);csvout('review.csv',review);csvout('conflicts.csv',conflicts)
    save('blacklist_gate.json',{'blacklisted_articles':len(blacklist_set),'flagged_active_cards':len(flagged)})
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name,rows in [('К согласованию',[r for r in review if r['status']=='PROPOSED']),('Есть источник, отложены',[r for r in review if r['status']=='HOLD' and r['source_rows']]),('Расхождения',conflicts),('Все пустые карточки',review)]:
        ws=wb.create_sheet(name)
        if not rows:ws.append(['Нет позиций']);continue
        ws.append(list(rows[0]))
        for r in rows:ws.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in r.values()])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476B')
        for i in range(1,len(rows[0])+1):ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=24
    wb.save(OUT/'first_wave_proposal.xlsx')
    print(json.dumps({'cards':len(ready),'stock_cards':sum(r['stock_positive'] for r in ready),'suppliers':summary},ensure_ascii=False))

if __name__=='__main__':main()
