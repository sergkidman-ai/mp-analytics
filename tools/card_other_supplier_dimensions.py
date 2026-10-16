#!/usr/bin/env python3
# поток: card
"""Read saved fresh supplier/MS snapshots; prepare dimensions coverage without writes."""
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools.card_sprint_dimensions_prepare import FIELDS, number

OUT = ROOT / 'docs/reports/card_other_supplier_dimensions_2026-10-14'
LABELS = {'rapid': 'Рапид', 'bulat': 'Булат', 'profiline': 'Профилайн', 'easy_print': 'Изипринт'}


def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def key(value):
    return re.sub(r'\s+', '', str(value or '')).upper()


def triplet(raw):
    m = re.fullmatch(r'\s*(\d+(?:[.,]\d+)?)\s*[xх×*]\s*(\d+(?:[.,]\d+)?)\s*[xх×*]\s*(\d+(?:[.,]\d+)?)\s*(?:мм)?\s*', str(raw), re.I)
    return bool(m and all(number(v) is not None for v in m.groups()))


def write_csv(name, rows):
    if rows:
        with (OUT / name).open('w', encoding='utf-8-sig', newline='') as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=';')
            w.writeheader()
            w.writerows(rows)


def main(all_rapid=False):
    source_folder = ROOT / 'docs/reports/card_other_supplier_dimensions_2026-10-14' if all_rapid else OUT
    indexes = {s: defaultdict(dict) for s in LABELS}
    rapid = json.loads((source_folder / 'rapid.json').read_text())
    for o in rapid['price']:
        raw = [o.get(n) for n in ['GabarityDlina', 'GabarityShirina', 'GabarityVisota']]
        nums = [number(v) for v in raw]
        cm = [v * 100 for v in nums] if all(v is not None for v in nums) else None
        valid = cm is not None and all(v <= 150 for v in cm)
        r = {'article': o.get('CodeID'), 'name': o.get('Name'), 'raw': ' × '.join(str(v) for v in raw),
             'unit': 'м', 'cm': [format(v.normalize(), 'f') for v in cm] if valid else [],
             'has_three': valid, 'source_id': str(o['ID']), 'url': 'https://b2b-rapid1.ru/api/export.php',
             'dimension_text': '', 'source_unit_confirmed': True}
        if key(r['article']):
            indexes['rapid'][key(r['article'])][r['source_id']] = r
    profiline = json.loads((source_folder / 'profiline.json').read_text())
    for o in profiline['products'].values():
        pp = {p['name']: p.get('value') for p in o.get('params', [])}
        raw = pp.get('Габариты в упаковке', '')
        r = {'article': o.get('vendorCode'), 'name': o['name'], 'raw': raw, 'unit': 'не подписана в поле API',
             'cm': [], 'has_three': triplet(raw), 'source_id': o['code'],
             'url': 'https://rm-profiline.online/json/hannover/client763/wms5/brand0/',
             'dimension_text': '', 'source_unit_confirmed': False}
        if key(r['article']):
            indexes['profiline'][key(r['article'])][r['source_id']] = r
    bulat = ET.fromstring((source_folder / 'bulat.xml').read_bytes())
    for o in bulat.findall('.//offer'):
        desc = o.findtext('description') or ''
        # Presence of words is evidence for manual review, not a package dimension.
        mentions = [desc[max(0, m.start()-40):m.start()+200] for m in list(re.finditer(
            r'габарит|размер.{0,20}упаков|упаков.{0,20}размер|ширина|высота|длина', desc, re.I))[:5]]
        r = {'article': o.get('id'), 'name': o.findtext('name'), 'raw': '', 'unit': '', 'cm': [],
             'has_three': False, 'source_id': o.get('id'), 'url': o.findtext('url'),
             'dimension_text': ' | '.join(mentions), 'source_unit_confirmed': False}
        for a in {o.get('id'), *(v.text for v in o.findall('vendorCode'))}:
            if key(a):
                indexes['bulat'][key(a)][r['source_id']] = r

    excluded = {}
    if all_rapid:
        from prices import blacklist
        black = blacklist.load_set()
        active = json.loads((OUT / 'ms_active_snapshot.json').read_text())
        scopes = {'rapid': {c['id']: c for c in active}}
        for c in active:
            reason = ''
            if blacklist.norm(c.get('article')) in black:
                reason = 'BLACKLIST'
            elif re.search(r'\bбрак\b', c.get('name', ''), re.I):
                reason = 'DAMAGED_CARD'
            elif any('стоп' in a.get('name','').lower() and str(a.get('value')).strip().lower() in ('true','да','1') for a in c.get('attributes',[])):
                reason = 'STOP_LIST'
            if reason:
                excluded[c['id']] = reason
        save('excluded_cards.json', excluded)
    else:
        scopes = {s: {} for s in LABELS}
        for doc in json.loads((OUT / 'current_positions_expanded.json').read_text()):
            for p in doc['positions']:
                c = p['assortment']
                scopes[doc['profile_key']][c['id']] = c
    metadata = json.loads((ROOT / 'docs/reports/card_sprint_dimensions_prepare_2026-10-14/ms_attribute_metadata.json').read_text())['rows']
    attrs = {n: next(a for a in metadata if a['name'] == n) for n in FIELDS}
    review, draft, summary = [], [], {}
    for supplier, cards in scopes.items():
        for c in cards.values():
            if c.get('archived') or c['meta']['type'] != 'product':
                continue
            old = {a['name']: a.get('value') for a in c.get('attributes', [])}
            old_raw = [old.get(n) for n in FIELDS]
            empty = [v is None or str(v).strip() == '' for v in old_raw]
            found = list(indexes[supplier].get(key(c.get('article')), {}).values())
            if len(found) > 1:
                exact = [r for r in found if str(r['article']).strip() == str(c.get('article') or '').strip()]
                if len(exact) == 1:
                    found = exact
            s = found[0] if len(found) == 1 else {}
            target = [Decimal(v) for v in s.get('cm', [])]
            status = 'SOURCE_UNMATCHED'
            if supplier == 'easy_print':
                status = 'CATALOGUE_UNAVAILABLE'
            elif len(found) == 1:
                if supplier == 'profiline':
                    status = 'PROFILINE_DIMENSIONS_NOT_APPROVED' if s['has_three'] else 'NO_PACK_DIMENSIONS'
                elif not s['has_three']:
                    if supplier == 'bulat' and c.get('article') in ('JFEP0ERC31020', 'MFBRPTH110010'):
                        # Manually checked source descriptions: ribbon width and
                        # length, not three dimensions of the product's package.
                        status = 'RIBBON_SIZE_NOT_PACKAGE_DIMENSIONS'
                    else:
                        status = 'DIMENSION_TEXT_REVIEW' if s['dimension_text'] else 'NO_THREE_DIMENSIONS'
                else:
                    values = [number(v) for v in old_raw]
                    if supplier == 'rapid' and key(c.get('article')) == 'SF-T9451BK':
                        # Source literally says 2.9 x 1.3 x 4 cm at 0.1 kg.
                        # Hold for verification; never infer a corrected size.
                        status = 'SUPPLIER_DIMENSIONS_SUSPECT_REVIEW'
                    elif values == target:
                        status = 'ALREADY_EQUAL'
                    elif all(v is not None for v in values) and sorted(values) == sorted(target):
                        status = 'SAME_SIDES_DIFFERENT_ORDER'
                    elif all(empty):
                        status = 'FILL_EMPTY'
                    elif any(empty) and all(e or v == t for e, v, t in zip(empty, values, target)):
                        status = 'FILL_PARTIAL_EMPTY'
                    else:
                        status = 'EXISTING_VALUES_REVIEW'
            if status in ('FILL_EMPTY', 'FILL_PARTIAL_EMPTY') and c['id'] in excluded:
                status = 'EXCLUDED_' + excluded[c['id']]
            r = {'supplier': supplier, 'id': c['id'], 'external_code': c.get('externalCode', ''),
                 'article': c.get('article', ''), 'ms_name': c['name'], 'source_match': len(found) == 1,
                 'supplier_name': s.get('name', ''), 'source_id': s.get('source_id', ''),
                 'source_url': s.get('url', ''), 'raw_dimensions': s.get('raw', ''), 'source_unit': s.get('unit', ''),
                 'source_three_dimensions': s.get('has_three', False), 'empty_ms_fields': sum(empty),
                 'old_length_cm': old_raw[0], 'old_width_cm': old_raw[1], 'old_height_cm': old_raw[2],
                 'proposed_length_cm': s.get('cm', ['','',''])[0] if s.get('cm') else '',
                 'proposed_width_cm': s.get('cm', ['','',''])[1] if s.get('cm') else '',
                 'proposed_height_cm': s.get('cm', ['','',''])[2] if s.get('cm') else '',
                 'dimension_text': s.get('dimension_text', ''), 'status': status}
            review.append(r)
            if status in ('FILL_EMPTY', 'FILL_PARTIAL_EMPTY'):
                updates = {attrs[n]['id']: {'id': attrs[n]['id'], 'meta': attrs[n]['meta'], 'value': format(v.normalize(), 'f')}
                           for n, v, e in zip(FIELDS, target, empty) if e}
                before = c.get('attributes', [])
                merged = [{**a, 'value': updates[a['id']]['value']} if a.get('id') in updates else a for a in before]
                present = {a.get('id') for a in before}
                merged.extend(a for aid, a in updates.items() if aid not in present)
                draft.append({'id': c['id'], 'supplier': supplier, 'external_code': c.get('externalCode'),
                              'article': c.get('article'), 'old_attributes': before, 'payload_preview': {'attributes': merged},
                              'changed_attribute_ids': list(updates), 'source_id': s['source_id'], 'source_raw': s['raw'],
                              'source_unit': s['unit'], 'authorized_to_write': False})
        selected = [r for r in review if r['supplier'] == supplier]
        summary[supplier] = {('active_ms_cards' if all_rapid else 'current_receipt_cards'): len(selected), 'matched_source_cards': sum(r['source_match'] for r in selected),
                             'cards_with_all_three_ms_fields_empty': sum(r['empty_ms_fields'] == 3 for r in selected),
                             'cards_with_any_ms_dimension_empty': sum(r['empty_ms_fields'] > 0 for r in selected),
                             'source_three_dimensions': sum(r['source_three_dimensions'] for r in selected),
                             'source_coverage_pct': round(sum(r['source_three_dimensions'] for r in selected)/len(selected)*100, 2) if selected and supplier != 'easy_print' else None,
                             'source_dimensions_and_empty_ms': sum(r['source_three_dimensions'] and r['empty_ms_fields'] > 0 for r in selected),
                             'ready_to_fill': sum(r['status'] in ('FILL_EMPTY', 'FILL_PARTIAL_EMPTY') for r in selected),
                             'statuses': dict(Counter(r['status'] for r in selected))}
    save('review.json', review)
    save('upload_fill_empty_preview.json', draft)
    save('summary.json', {'at_utc': datetime.now(timezone.utc).isoformat(), 'scope': 'all_active_rapid' if all_rapid else 'current_receipts', 'suppliers': summary,
                          'source_dates': {'rapid': rapid.get('date'), 'bulat': bulat.get('date'), 'profiline': profiline.get('date')},
                          'external_writes': 0, 'database_writes': 0})
    write_csv('review.csv', review)
    write_csv('ready_to_fill.csv', [r for r in review if r['status'] in ('FILL_EMPTY','FILL_PARTIAL_EMPTY')])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Итог'
    ws.append(['Поставщик', 'Активных карточек МС' if all_rapid else 'Карточек в оприходованиях', 'Есть три размера в источнике', 'Покрытие, %',
               'МС: есть пустые поля размеров', 'Размеры источника + пустые МС', 'Подготовлено к заполнению'])
    for supplier, s in summary.items():
        ws.append([LABELS[supplier], s['active_ms_cards' if all_rapid else 'current_receipt_cards'], s['source_three_dimensions'], s['source_coverage_pct'],
                   s['cards_with_any_ms_dimension_empty'], s['source_dimensions_and_empty_ms'], s['ready_to_fill']])
    ws.append(['Примечание', 'Профилайн исключён действующим правилом; у Изипринта каталог не получен. Размеры из объёма не рассчитывались.'])
    for title, rows in [('К заполнению', [r for r in review if r['status'] in ('FILL_EMPTY','FILL_PARTIAL_EMPTY')]),
                        ('Профилайн отдельно', [r for r in review if r['supplier']=='profiline' and r['source_three_dimensions'] and r['empty_ms_fields']>0]),
                        ('Все карточки', review)]:
        sheet = wb.create_sheet(title)
        sheet.append(list(review[0]))
        for r in rows:
            sheet.append(list(r.values()))
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='305496')
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(16, max(len(str(c.value or '')) for c in col[:20])+2))
    wb.save(OUT / 'supplier_dimensions_review.xlsx')
    print(json.dumps(summary, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
