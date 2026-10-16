# поток: card
"""Read-only proposal from each supplier's own source; no cross-supplier comparisons."""
import csv
import json
import re
import sys
import hashlib
from collections import Counter,defaultdict
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from pathlib import Path
import openpyxl
from openpyxl.styles import Font,PatternFill

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from prices import blacklist
AUDIT=ROOT/'docs/reports/card_all_supplier_dimensions_2026-10-15'
OUT=ROOT/'docs/reports/card_dimensions_own_source_plan_2026-10-15'
GROUPS={'Кактус':'Кактус','Сакура':'Сакура','Изипринт':'Изипринт','Рапид / Суперфайн':'Рапид'}
# Individually reviewed defects of the supplier's own source, never external-code comparisons.
SOURCE_HOLDS={
    ('Кактус','CS-CLI521BK/M/C/Y'):'SET_WITH_COMPONENT_SIZE: 4-cartridge set has 85.5x37x15.5mm; package unresolved',
    ('Кактус','CS-CLI426BK/M/C/Y'):'SET_WITH_COMPONENT_SIZE: 4-cartridge set has 85.5x37x15.5mm; package unresolved',
    ('Сакура','SA006R01179'):'OWN_SOURCE_CONTRADICTION: 43x75x75cm gives 241.875L but same row volume is 2.419L',
    ('Сакура','SACF237A'):'OBVIOUS_PLACEHOLDER: 10x10x10cm for full laser cartridge; same source volume 12.863L',
    ('Сакура','SADR2510'):'SUSPECT_DECIMAL: 31x90x21.5cm for individual drum; 90cm side needs supplier confirmation',
    ('Изипринт','TC-HW1106AL'):'SUSPECT_MASTER_CARTON: 52x33.5x50cm for single HP107 cartridge, individual package unconfirmed',
}
for code in ['SAW2000A','SAW2001A','SAW2002A','SAW2003A','SAW2000X','SAW2001X','SAW2002X','SAW2003X']:
    SOURCE_HOLDS[('Сакура',code)]='OWN_SOURCE_CONTRADICTION: identical 39cm sides give 59.319L, source volume is approximately 10L'
for code in ['SIC13T741200','SIC13T741300','SIC13T741400','SIC13T741X00']:
    SOURCE_HOLDS[('Сакура',code)]='OWN_SOURCE_CONTRADICTION: 50.5x33x50cm gives 83.325L, source volume 2.931L for 1L ink'

def loadcsv(path):
    with path.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def num(v):
    try:
        d=Decimal(str(v).strip().replace(',','.'));return d if d.is_finite() and d>0 else None
    except InvalidOperation:return None
def key(v):return str(v or '').strip().upper()
def fmt(ds):return [format(d.normalize(),'f') for d in ds]
def save(name,d):(OUT/name).write_text(json.dumps(d,ensure_ascii=False,indent=2))
def csvout(name,rows):
    if not rows:return
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

def main():
    OUT.mkdir(exist_ok=True)
    cards=loadcsv(AUDIT/'products_review.csv');src=loadcsv(AUDIT/'source_rows.csv')
    index=defaultdict(list)
    for s in src:
        if s['supplier'] not in GROUPS.values() or s['three_positive_values']!='True' or s['units'] not in ['мм','см','м']:continue
        if s['supplier']=='Рапид' and '2026-10-14/rapid.json' not in s['file']:continue
        vals=[num(v) for v in json.loads(s['raw_dimensions'])]
        if s['file'].endswith('.xlsx'):vals=[Decimal(format(v,'.15g')) for v in vals]
        factor={'мм':Decimal('.1'),'см':Decimal(1),'м':Decimal(100)}[s['units']]
        index[(s['supplier'],key(s['article']))].append(dict(s,cm=[v*factor for v in vals]))
    blocked={};black=blacklist.load_set()
    for p in json.loads((AUDIT/'products.json').read_text()):
        reasons=[]
        if blacklist.norm(p.get('article')) in black:reasons.append('BLACKLIST')
        if re.search(r'\bбрак\b',p.get('name',''),re.I):reasons.append('DAMAGED_CARD')
        if blacklist.in_damaged(p.get('name')):reasons.append('DAMAGED_PACKAGE')
        if any('стоп' in a.get('name','').lower() and str(a.get('value')).lower() in ['true','да','1'] for a in p.get('attributes',[])):reasons.append('STOP_LIST')
        if reasons:blocked[p['id']]=reasons
    review=[];fills=[];replacements=[];outliers=[]
    for c in cards:
        if c['group'] not in GROUPS:continue
        supplier=GROUPS[c['group']]
        ss=index.get((supplier,key(c['article'])),[])
        signatures={tuple(s['cm']) for s in ss}
        reasons=list(blocked.get(c['ms_id'],[]))
        if not ss:reasons.append('NO_EXPLICIT_DIRECT_SOURCE')
        elif len(signatures)!=1:reasons.append('AMBIGUOUS_OWN_SOURCE')
        ds=ss[0]['cm'] if len(signatures)==1 else None
        if supplier=='Рапид' and not key(c['article']).startswith(('SF','OEM')):reasons.append('RAPID_ORIGINALS_PREVIOUSLY_EXCLUDED')
        if supplier=='Рапид' and key(c['article'])=='SF-T9451BK':reasons.append('KNOWN_SUSPECT_3065')
        if (supplier,key(c['article'])) in SOURCE_HOLDS:reasons.append('REVIEWED_OWN_SOURCE_DEFECT')
        current=[num(c[k]) for k in ['length_cm','width_cm','height_cm']]
        diagnostics=[]
        if ds:
            if min(ds)<1:diagnostics.append('SIDE_UNDER_1CM')
            if max(ds)>100:diagnostics.append('SIDE_OVER_100CM')
            if ds[0]*ds[1]*ds[2]>25000:diagnostics.append('BOX_OVER_25L')
            if re.search(r'\b(?:набор|комплект)\w*\b',c['name'],re.I):diagnostics.append('SUPPLIER_SET_PACKAGE')
            if max(ds)>150 and re.search(r'картридж|к-ж|тонер|фотобарабан|драм',c['name'],re.I):
                reasons.append('IMPLAUSIBLE_CONSUMABLE_SIZE_OVER_150CM')
            if min(ds)<1 and re.search(r'картридж|к-ж|фотобарабан|драм',c['name'],re.I):
                reasons.append('IMPLAUSIBLE_CARTRIDGE_SIDE_UNDER_1CM')
        if reasons:status='HOLD'
        elif all(current) and sorted(current)==sorted(ds):status='NO_CHANGE_SAME_BOX'
        elif all(current) and sorted(current)==sorted(d.quantize(Decimal('.01'),rounding=ROUND_HALF_UP) for d in ds):
            status='NO_CHANGE_FORMAT_PRECISION'
        elif c['all_three_filled']=='True':status='CANDIDATE_REPLACE'
        else:status='CANDIDATE_FILL'
        row={'supplier':c['group'],'ms_id':c['ms_id'],'supplier_id':c['supplier_id'],'external_code':c['external_code'],
            'article':c['article'],'name':c['name'],'stock_positive':c['any_target_stock']=='True',
            'current_cm':' × '.join(fmt(current)) if all(current) else '',
            'source_cm':' × '.join(fmt(ds)) if ds else '',
            'status':status,'reasons':';'.join(reasons),'diagnostic_flags':';'.join(diagnostics),
            'own_source_defect_note':SOURCE_HOLDS.get((supplier,key(c['article'])),''),
            'source_rows':len(ss),'source_file':ss[0]['file'] if ss else '',
            'source_sheet':ss[0]['sheet'] if ss else '', 'source_row':ss[0]['row_or_item'] if ss else ''}
        review.append(row)
        if diagnostics and ds:outliers.append(row)
        if status.startswith('CANDIDATE_'):
            proposed=dict(row,dimensions_cm=fmt(ds),source_pointers=[{'file':s['file'],'sheet':s['sheet'],'row':s['row_or_item']} for s in ss])
            (fills if status=='CANDIDATE_FILL' else replacements).append(proposed)
    summary={}
    for g in GROUPS:
        rr=[r for r in review if r['supplier']==g]
        summary[g]={'active_cards':len(rr),'with_source':sum(r['source_rows']>0 for r in rr),
            'fill':sum(r['status']=='CANDIDATE_FILL' for r in rr),'fill_stock':sum(r['status']=='CANDIDATE_FILL' and r['stock_positive'] for r in rr),
            'replace':sum(r['status']=='CANDIDATE_REPLACE' for r in rr),'replace_stock':sum(r['status']=='CANDIDATE_REPLACE' and r['stock_positive'] for r in rr),
            'same_box':sum(r['status']=='NO_CHANGE_SAME_BOX' for r in rr),
            'same_at_stored_precision':sum(r['status']=='NO_CHANGE_FORMAT_PRECISION' for r in rr),
            'hold':sum(r['status']=='HOLD' for r in rr),'reason_counts':dict(Counter(x for r in rr for x in r['reasons'].split(';') if x))}
    save('summary.json',{'status':'PROPOSAL_NOT_AUTHORIZED','policy':'own supplier source only; NO comparison by external code and NO percentage threshold',
        'suppliers':summary,'fill_cards':len(fills),'replace_cards':len(replacements),'all_change_cards':len(fills)+len(replacements),
        'scope_question':'replacement policy asked; no material replacement candidates remain in this wave',
        'outlier_review_pending':False,'writes':0,'scheduled':False})
    save('candidate_fill.json',fills);save('candidate_replace.json',replacements)
    save('own_source_defects.json',[{'supplier':k[0],'article':k[1],'reason':v} for k,v in SOURCE_HOLDS.items()])
    csvout('review.csv',review);csvout('source_outliers.csv',outliers)
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name,rr in [('Пустые',[r for r in review if r['status']=='CANDIDATE_FILL']),('Заполненные отличаются',[r for r in review if r['status']=='CANDIDATE_REPLACE']),('Проверка исходника',outliers),('Отложено с источником',[r for r in review if r['status']=='HOLD' and r['source_rows']]),('Все карточки',review)]:
        ws=wb.create_sheet(name)
        if not rr:ws.append(['Нет позиций']);continue
        ws.append(list(rr[0]))
        for r in rr:ws.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in r.values()])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476B')
        for i in range(1,len(rr[0])+1):ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=24
    wb.save(OUT/'own_supplier_source_proposal.xlsx')
    print(json.dumps({'fill':len(fills),'replace':len(replacements),'outlier_rows':len(outliers),'suppliers':summary},ensure_ascii=False))

if __name__=='__main__':main()
