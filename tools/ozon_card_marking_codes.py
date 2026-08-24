"""tools/ozon_card_marking_codes.py — поток: card.

Внешние коды МС для карточек Ozon, которым нужен код маркировки (класс C, acc2).

Только чтение. Мост Ozon -> МойСклад: offer_id аккаунта Дисквэр устроен как
<4-значный внешний код МС> + 8 случайных символов (13 558 из 13 616 карточек acc2 длиной
ровно 12; сверка по названиям модели совпала на выборке). Штрихкодами мост не строится:
у этих карточек barcodes пуст и на площадке, и в v4-атрибутах.
"""
import os, csv, collections
import psycopg2
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

load_dotenv("/opt/mp-analytics/.env")
SRC = "/opt/mp-analytics/docs/reports/ozon_card_content_errors.csv"
OUT = "/opt/mp-analytics/docs/reports/ozon_marking_codes_2026-08-23.xlsx"

rows = [r for r in csv.DictReader(open(SRC, encoding="utf-8"))
        if "маркиров" in (r["attribute_name"] or "")]
print("карточек в задании:", len(rows))
print("длины offer_id:", collections.Counter(len(r["offer_id"]) for r in rows).most_common())

codes = sorted({r["offer_id"][:4] for r in rows})
c = psycopg2.connect(os.environ["DATABASE_URL"]); cur = c.cursor()
cur.execute("""select payload->>'externalCode', payload->>'name', payload->>'code',
                      payload->>'article', payload->>'trackingType', payload->>'archived'
               from raw_moysklad_product where payload->>'externalCode' = any(%s)""", (codes,))
ms = collections.defaultdict(list)
for ext, name, code, art, trk, arch in cur.fetchall():
    ms[ext].append({"name": name or "", "code": code or "", "article": art or "",
                    "tracking": trk or "", "archived": arch == "true"})
print("внешних кодов в задании:", len(codes), "| найдено в МС:", len(ms))

out = []
for r in rows:
    ext = r["offer_id"][:4]
    variants = [v for v in ms.get(ext, []) if not v["archived"]] or ms.get(ext, [])
    v = variants[0] if variants else {}
    out.append({
        "Внешний код МС": ext,
        "Есть в МС": "да" if variants else "НЕТ",
        "Наименование МС": v.get("name", ""),
        "Код МС": v.get("code", ""),
        "Тип маркировки в МС": v.get("tracking", ""),
        "Позиций поставщиков под кодом": len(ms.get(ext, [])),
        "offer_id Ozon": r["offer_id"],
        "product_id Ozon": r["product_id"],
        "Состояние карточки": r["status_descr"],
        "Наименование на Ozon": r["name"],
        "Что требует Ozon": "заполнить атрибут «Нужен код маркировки»",
    })
out.sort(key=lambda x: (x["Есть в МС"] == "да", x["Внешний код МС"]))
nomatch = sum(1 for x in out if x["Есть в МС"] == "НЕТ")
print("сшито с МС:", len(out) - nomatch, "| не нашлось:", nomatch)
print("уникальных внешних кодов в файле:", len({x["Внешний код МС"] for x in out}))

wb = Workbook(); ws = wb.active; ws.title = "Нужен код маркировки"
cols = list(out[0].keys())
ws.append(cols)
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
for r in out:
    ws.append([r[k] for k in cols])
widths = {"Внешний код МС": 16, "Есть в МС": 11, "Наименование МС": 60, "Код МС": 12,
          "Тип маркировки в МС": 20, "Позиций поставщиков под кодом": 14,
          "offer_id Ozon": 16, "product_id Ozon": 16, "Состояние карточки": 18,
          "Наименование на Ozon": 60, "Что требует Ozon": 38}
for i, k in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 18)
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(out) + 1}"
wb.save(OUT)
print("файл:", OUT)
