#!/usr/bin/env python3
# поток: card
"""Read-only Sprint dimensions collection and MS import preview. No MS writes."""
import argparse
import csv
import json
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import requests
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from core import ms_api  # loads credentials from .env
from tools.card_solutionsprint_scraper import login, parse_card
from tools.card_solutionsprint_description_audit import fetch, listing

OUT = ROOT / 'docs/reports/card_sprint_dimensions_prepare_2026-10-14'
FIELDS = ['Длина, см.', 'Ширина, см.', 'Высота, см.']


def save(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def number(value):
    try:
        d = Decimal(str(value).strip().replace(',', '.'))
        return d if d.is_finite() and d > 0 else None
    except InvalidOperation:
        return None


def dimensions(raw):
    m = re.fullmatch(r'\s*(\d+(?:[.,]\d+)?)\s*[xх×*]\s*(\d+(?:[.,]\d+)?)\s*[xх×*]\s*(\d+(?:[.,]\d+)?)\s*(?:мм)?\s*', raw, re.I)
    if not m:
        return None
    values = [number(s) for s in m.groups()]
    return [v / 10 for v in values] if all(v is not None for v in values) else None


def code(card):
    m = re.fullmatch(r'SP(\d+)(?:_MSK)?', card.get('article', ''), re.I)
    return m[1] if m else None


def collect(cards):
    cache = OUT / 'supplier_cards'
    cache.mkdir(exist_ok=True)
    index = defaultdict(dict)
    old = ROOT / 'docs/reports/card_solutionsprint_current_review_2026-10-12/source_listings.json'
    for r in json.loads(old.read_text()):
        index[r['supplier_code']][r['url']] = r
    local = threading.local()
    lock = threading.Lock()
    next_start = [0.0]

    def doc(url, params=None):
        if not hasattr(local, 'session'):
            local.session = requests.Session()
            login(local.session)
        with lock:
            time.sleep(max(0, next_start[0] - time.monotonic()))
            next_start[0] = time.monotonic() + .5
        return fetch(local.session, url, params)

    def one(c):
        path = cache / (c + '.json')
        if path.exists():
            return json.loads(path.read_text())
        try:
            found = list(index.get(c, {}).values())
            if len(found) != 1:
                found = [r for r in listing(doc('/search/index.php', {'q': c, 's': 'Поиск'})) if r['supplier_code'] == c]
            if len(found) != 1:
                result = {'supplier_code': c, 'status': 'NO_UNIQUE_URL'}
            else:
                page = doc(found[0]['url'])
                result = parse_card(html.tostring(page), found[0]['url'], authenticated=True)
                if result['supplier_code'] != c:
                    raise ValueError('Supplier code mismatch')
                result['status'] = 'READ_OK'
        except (RuntimeError, ValueError) as e:
            result = {'supplier_code': c, 'status': 'READ_ERROR', 'error_type': type(e).__name__}
        result['collected_at_utc'] = datetime.now(timezone.utc).isoformat()
        path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
        return result

    wanted = sorted({code(c) for c in cards if code(c)})
    results = []
    with ThreadPoolExecutor(max_workers=4) as pool:
        for task in as_completed([pool.submit(one, c) for c in wanted]):
            results.append(task.result())
            if len(results) % 50 == 0 or len(results) == len(wanted):
                progress = {'done': len(results), 'total': len(wanted), 'statuses': dict(Counter(r['status'] for r in results))}
                save('progress.json', progress)
                print(json.dumps(progress), flush=True)
    save('supplier_cards.json', results)
    return {r['supplier_code']: r for r in results}


def csv_save(name, rows):
    if not rows:
        return
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=';')
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--scope', choices=['receipts', 'all'], default='receipts')
    args = ap.parse_args()
    raw = json.loads((OUT / ('receipt_cards.json' if args.scope == 'receipts' else 'ms_active_snapshot.json')).read_text())
    cards = [c for c in raw if not c.get('archived')]
    metadata = json.loads((OUT / 'ms_attribute_metadata.json').read_text())['rows']
    attrs = {n: [a for a in metadata if a['name'] == n] for n in FIELDS}
    assert all(len(a) == 1 and a[0]['type'] == 'string' for a in attrs.values())
    source = collect(cards)
    review, previews = [], []
    for c in cards:
        s = source.get(code(c), {})
        props = s.get('properties', {})
        raw_dim = props.get('Габариты, мм', '')
        new = dimensions(raw_dim)
        by_name = {a['name']: a.get('value') for a in c.get('attributes', [])}
        old_raw = [by_name.get(n, '') for n in FIELDS]
        old = [number(v) for v in old_raw]
        if not code(c):
            status = 'ARTICLE_UNMATCHED'
        elif s.get('status') != 'READ_OK':
            status = s.get('status', 'SOURCE_MISSING')
        elif new is None:
            status = 'SUPPLIER_DIMENSIONS_INVALID_OR_MISSING'
        elif 'лазер' in props.get('Вид товара', '').lower() and min(new) < 1:
            # A review flag, never a rule to infer or repair a physical dimension.
            status = 'SUPPLIER_DIMENSIONS_SUSPECT_REVIEW'
        elif old == new:
            status = 'ALREADY_EQUAL'
        elif all(v is not None for v in old) and sorted(old) == sorted(new):
            status = 'SAME_SIDES_DIFFERENT_ORDER'
        elif all(v is None or str(v).strip() == '' for v in old_raw):
            status = 'FILL_ALL_THREE'
        elif any(v is None for v in old):
            status = 'PARTIAL_FIELDS_REVIEW'
        else:
            status = 'DIFFERENT_DIMENSIONS_REVIEW'
        row = {'id': c['id'], 'external_code': c.get('externalCode', ''), 'article': c.get('article', ''),
               'ms_name': c['name'], 'supplier_code': code(c), 'supplier_name': s.get('name', ''),
               'url': s.get('url', ''), 'supplier_dimensions_mm': raw_dim,
               'old_length_cm': old_raw[0], 'old_width_cm': old_raw[1], 'old_height_cm': old_raw[2],
               'proposed_length_cm': format(new[0].normalize(), 'f') if new else '',
               'proposed_width_cm': format(new[1].normalize(), 'f') if new else '',
               'proposed_height_cm': format(new[2].normalize(), 'f') if new else '', 'status': status,
               'source_at_utc': s.get('collected_at_utc', '')}
        review.append(row)
        if new and status not in ('ALREADY_EQUAL', 'SAME_SIDES_DIFFERENT_ORDER', 'SUPPLIER_DIMENSIONS_SUSPECT_REVIEW'):
            previews.append({'id': c['id'], 'article': c.get('article'), 'external_code': c.get('externalCode'),
                             'status': status, 'old_dimensions': dict(zip(FIELDS, old_raw)),
                             'attributes': [{'meta': attrs[n][0]['meta'], 'id': attrs[n][0]['id'], 'value': format(v.normalize(), 'f')}
                                            for n, v in zip(FIELDS, new)],
                             'source_url': s['url'], 'raw_mm': raw_dim,
                             'axis_mapping': 'First/second/third supplier value -> length/width/height, preserving source order; supplier does not label axes',
                             'authorized_to_write': False})
    save('review.json', review)
    save('preview_dimensions.json', previews)
    # Preparation only: preserve unrelated attributes and restrict the import
    # draft to genuinely empty fields. Re-read before any separately authorized write.
    by_id = {c['id']: c for c in cards}
    fills = []
    for preview in previews:
        if preview['status'] != 'FILL_ALL_THREE':
            continue
        card = by_id[preview['id']]
        updates = {a['id']: a for a in preview['attributes']}
        old_attributes = card.get('attributes', [])
        merged = [{**a, 'value': updates[a['id']]['value']} if a.get('id') in updates else a
                  for a in old_attributes]
        present = {a.get('id') for a in old_attributes}
        merged.extend(a for a in preview['attributes'] if a['id'] not in present)
        fills.append({'id': card['id'], 'article': card.get('article'),
                      'source_url': preview['source_url'], 'raw_mm': preview['raw_mm'],
                      'old_attributes': old_attributes, 'proposed_payload': {'attributes': merged},
                      'authorized_to_write': False, 'requires_live_read_before_write': True})
    save('upload_fill_empty_preview.json', fills)
    csv_save('review.csv', review)
    csv_save('candidates.csv', [r for r in review if r['status'] in ('FILL_ALL_THREE', 'PARTIAL_FIELDS_REVIEW', 'DIFFERENT_DIMENSIONS_REVIEW')])
    counts = Counter(r['status'] for r in review)
    covered = sum(dimensions(r['supplier_dimensions_mm']) is not None for r in review)
    snapshot_path = OUT / ('receipt_cards.json' if args.scope == 'receipts' else 'ms_active_snapshot.json')
    summary = {'at_utc': datetime.now(timezone.utc).isoformat(),
               'ms_snapshot_at_utc': datetime.fromtimestamp(snapshot_path.stat().st_mtime, timezone.utc).isoformat(),
               'scope': args.scope, 'active_ms_cards': len(cards),
               'archived_excluded': len(raw)-len(cards), 'unique_supplier_codes': len(source),
               'source_statuses': dict(Counter(s['status'] for s in source.values())),
               'cards_with_valid_supplier_dimensions': covered, 'coverage_pct': round(covered/len(cards)*100, 2),
               'cards_without_suspect_dimension_flag': covered-counts['SUPPLIER_DIMENSIONS_SUSPECT_REVIEW'],
               'coverage_without_suspect_pct': round((covered-counts['SUPPLIER_DIMENSIONS_SUSPECT_REVIEW'])/len(cards)*100, 2),
               'statuses': dict(counts), 'preview_rows': len(previews), 'fill_empty_draft_rows': len(fills),
               'ms_dimension_fields': {n: attrs[n][0]['id'] for n in FIELDS},
               'supplier_axes_explicitly_labelled': False,
               'axis_mapping': 'Supplier first/second/third -> MS length/width/height; no sorting',
               'axis_mapping_matches_existing_cards': counts['ALREADY_EQUAL'], 'external_writes': 0,
               'missing_receipt_profiles': json.loads((OUT/'missing_receipt_profiles.json').read_text())}
    save('summary.json', summary)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Итог'
    for row in [
        ['Показатель', 'Значение'], ['Область', args.scope], ['Активных карточек МС', len(cards)],
        ['Есть три положительных размера поставщика', covered], ['Покрытие, %', summary['coverage_pct']],
        ['Подозрительные размеры поставщика — проверить', counts['SUPPLIER_DIMENSIONS_SUSPECT_REVIEW']],
        ['Пустые поля — проект дозаполнения', len(fills)], ['Записей в МС', 0],
        ['Единицы', 'У поставщика мм; в МС см; деление на 10 без округления'],
        ['Порядок сторон', 'Сохраняется порядок поставщика; первый→Длина, второй→Ширина, третий→Высота. Сам поставщик оси не подписывает.'],
        ['Ограничение', 'Заполненные значения не перезаписывать. Перед записью повторно прочитать карточку.'],
        ['Нет текущих документов', ', '.join(summary['missing_receipt_profiles'])],
    ]:
        ws.append(row)
    headers = ['ID МС', 'Внешний код', 'Артикул МС', 'Название МС', 'Код Спринта', 'Название Спринта',
               'Ссылка', 'Габариты поставщика, мм', 'МС длина, см', 'МС ширина, см', 'МС высота, см',
               'Предлагаемая длина, см', 'Предлагаемая ширина, см', 'Предлагаемая высота, см', 'Статус', 'Дата источника UTC']
    for title, records in [('Все выбранные', review), ('Пустые поля', [r for r in review if r['status'] == 'FILL_ALL_THREE']),
                           ('На проверку', [r for r in review if r['status'] not in ('ALREADY_EQUAL', 'FILL_ALL_THREE')])]:
        sheet = wb.create_sheet(title)
        sheet.append(headers)
        for r in records:
            sheet.append(list(r.values()))
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='305496')
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(65, max(16, max(len(str(c.value or '')) for c in col[:30]) + 2))
    wb.save(OUT / 'sprint_dimensions_review.xlsx')
    print(json.dumps(summary, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
