# поток: prc
"""Чип в ТК: список на правку. База — согласие поставщиков; плюс наборы целиком.

Правило Сергея 21.08.2026: часть набора не может идти с чипом, а часть без. Поэтому если код
из подборки входит в набор (`set_cost`), в файл добавляются ВСЕ остальные компоненты этого
набора и сам набор — чтобы поле «чип» правилось по набору целиком, а не по одному цвету.
"""
import sys, ast, csv, json, collections, pathlib
sys.path.insert(0, "/opt/mp-analytics")
from core import db
from prices import features
RU = features.CHIP_NAMES

# --- наборы: код набора -> компоненты, компонент -> наборы
sets_of, comps_of = collections.defaultdict(list), {}
for r in db.query("select external_code ec, components c from set_cost"):
    v = r["c"]
    comps = [str(x).strip() for x in (ast.literal_eval(v) if isinstance(v, str) else v)]
    comps_of[r["ec"]] = comps
    for c in comps:
        sets_of[c].append(r["ec"])

# --- поставщики с остатком на «Удаленном» (свежий слепок)
sup = collections.defaultdict(list)
for r in db.query("""select external_code ec, supplier, name, stock from supplier_stock
                     where captured_at=(select max(captured_at) from supplier_stock)
                       and store='Удаленный склад' and stock>0"""):
    t = features.chip(r["name"])
    if t: sup[(r["ec"] or "").strip()].append((r["supplier"], r["name"], t, r["stock"]))

tk = {r["external_code"]: r for r in db.query("""select external_code, title, additional_title,
        united_title, color, resource, chip, consumable_type from prc_tc_model where gone_at is null""")}

# --- наши карточки МС по внешнему коду (слепок raw_moysklad_product)
# только нужные поля: тянуть 45 тыс. полных payload в память нельзя (OOM 21.08.2026)
ms = collections.defaultdict(list)
for r in db.query("""select payload->>'externalCode' ec, payload->>'code' code
                     from raw_moysklad_product
                     where coalesce((payload->>'archived')::bool, false) = false"""):
    ms[(r["ec"] or "").strip()].append(((r["code"] or ""), ""))

def consensus(ec):
    """-> (токен|None, сколько источников, статус)"""
    items = sup.get(ec, [])
    toks = collections.Counter(x[2] for x in items)
    if not toks: return None, 0, "поставщики про чип молчат"
    if len(toks) > 1:
        return None, sum(toks.values()), "разнобой: " + "; ".join(f"{RU[k]}×{v}" for k, v in toks.most_common())
    t, n = next(iter(toks)), sum(toks.values())
    return t, n, ("согласны " + str(n) + " поставщика" if n >= 2 else "источник один")

base = [(r["внешний код"] or "").strip() for r in csv.DictReader(
        open("/opt/mp-analytics/docs/reports/prc_tk_chip_to_fill_2026-08-21.csv",
             encoding="utf-8-sig"), delimiter=";")]
base = set(base)

# --- разворачиваем наборы
role = {ec: "согласие поставщиков" for ec in base}
touched = sorted({s for ec in base if ec in sets_of for s in sets_of[ec]})
for s in touched:
    role.setdefault(s, "НАБОР целиком")
    for c in comps_of.get(s, []):
        if c not in base: role.setdefault(c, "часть набора")

rows = []
for ec in sorted(role):
    t = tk.get(ec) or {}
    tok, n, status = consensus(ec)
    cur = RU[t.get("chip") or None] if t else "модели нет в ТК"
    # что ставить: своё согласие; иначе — по набору, если у набора единое мнение
    src_hint = ""
    put = RU[tok] if tok else ""
    if not put:
        mates = set()
        for s in sets_of.get(ec, []) + ([ec] if ec in comps_of else []):
            for c in comps_of.get(s, []) + [s]:
                mt = consensus(c)[0]
                if mt: mates.add(mt)
        if len(mates) == 1:
            put, src_hint = RU[next(iter(mates))], "по набору"
        else:
            put, src_hint = "?", "разобрать руками"
    items = sup.get(ec, [])
    rows.append({
        "внешний код": ec,
        "почему в списке": role[ec],
        "набор(ы)": ", ".join(sets_of.get(ec, [])) or (", ".join(comps_of.get(ec, [])[:8]) if ec in comps_of else ""),
        "модель ТК": t.get("title") or "",
        "доп. название": t.get("additional_title") or "",
        "серия": t.get("united_title") or "",
        "цвет": t.get("color") or "",
        "ресурс": t.get("resource") or "",
        "чип в ТК сейчас": cur,
        "поставить в ТК": put,
        "основание": src_hint or status,
        "источников": n,
        "как пишут поставщики": " | ".join(sorted({x[1][:46] for x in items})[:3]),
        "остаток Удаленный": sum(x[3] for x in items),
        "наши карточки МС": ", ".join(sorted({c for c, _ in ms.get(ec, [])})[:8]),
    })

st = collections.Counter(r["почему в списке"] for r in rows)
warn = [r for r in rows if r["чип в ТК сейчас"] not in ("не указан", "модели нет в ТК")
        and r["поставить в ТК"] not in ("?", r["чип в ТК сейчас"])]
print("строк в файле:", len(rows), "| задето наборов:", len(touched))
for k, v in st.most_common(): print(f"  {k:24} {v}")
print("  из них «?» (нечем решить):", sum(1 for r in rows if r["поставить в ТК"] == "?"))
print("  РАСХОЖДЕНИЕ с уже заполненным ТК:", len(warn))
for r in warn[:6]:
    print(f"    {r['внешний код']} {r['модель ТК'][:16]:16} ТК: {r['чип в ТК сейчас']:22} → {r['поставить в ТК']} ({r['основание']})")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
cols = list(rows[0].keys())
wb = Workbook(); ws = wb.active; ws.title = "чип — правка в ТК"
ws.append(cols)
for c in ws[1]: c.font = Font(bold=True)
fill = {"НАБОР целиком": PatternFill("solid", fgColor="FFF2CC"),
        "часть набора": PatternFill("solid", fgColor="E2EFDA")}
for r in rows:
    ws.append([r[c] for c in cols])
    if r["почему в списке"] in fill:
        for c in ws[ws.max_row]: c.fill = fill[r["почему в списке"]]
for i, w in enumerate((12, 22, 26, 18, 16, 22, 8, 10, 20, 20, 26, 11, 52, 18, 30), start=1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
p = pathlib.Path("/opt/mp-analytics/docs/reports/prc_tk_chip_to_fill_2026-08-21.xlsx")
wb.save(p)
with p.with_suffix(".csv").open("w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=cols, delimiter=";"); w.writeheader(); w.writerows(rows)
print("файл:", p)
