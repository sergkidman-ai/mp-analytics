#!/usr/bin/env python3
# поток: fin
"""Build a reviewed, hashed auto-action cap manifest; never send marketplace prices."""
import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
import zipfile

sys.path.insert(0, '/opt/mp-analytics')
from core import db
from core.wb_auto_entry import number
from ops.wb_auto_entry import atomic_json
from openpyxl import load_workbook

BRANDS = {'Цифровой квадрат': 'wb_acc1', 'Dsquare': 'wb_acc2'}


def build(dropbox, calendar, output):
    offers = {'wb_acc1': {}, 'wb_acc2': {}}
    sources = []
    seen = set()

    def offer(account, nm, cap, ids, source):
        if number(cap) <= 0:
            return
        promos = [p for p in calendar['accounts'][account] if p['id'] in ids]
        if len(promos) != len(ids) or len({(p['startDateTime'], p['endDateTime']) for p in promos}) != 1:
            raise ValueError('calendar generation mismatch')
        item = {'ids': ids, 'name': promos[0]['name'], 'start': promos[0]['startDateTime'],
                'end': promos[0]['endDateTime'], 'cap': str(number(cap)), 'source': source}
        prior = offers[account].setdefault(str(nm), [])
        if item not in prior:
            prior.append(item)

    def read(data, name, promo_id):
        digest = hashlib.sha256(data).hexdigest()
        if digest in seen:
            return
        seen.add(digest)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        sheet.reset_dimensions()  # WB's files incorrectly advertise dimension=A1.
        records = sheet.iter_rows(values_only=True)
        header = next(records)
        columns = {str(h).strip(): i for i, h in enumerate(header) if h}
        counts = {}
        for row in records:
            if not any(x is not None for x in row):
                continue
            account = BRANDS.get(str(row[columns['Бренд']]).strip())
            if not account:
                raise ValueError('unknown workbook brand')
            nm = int(number(row[columns['Артикул WB']]))
            cap = row[columns['Плановая цена для акции']]
            offer(account, nm, cap, [promo_id], name)
            counts[account] = counts.get(account, 0) + 1
        workbook.close()
        sources.append({'name': name, 'sha256': digest, 'rows': counts, 'promo_id': promo_id})

    xlsx = sorted(dropbox.glob('20261021_*Товары_для_исключения*11_11*.xlsx'))
    if len(xlsx) != 3:
        raise ValueError(f'expected three standalone 21.10 files, got {len(xlsx)}')
    for path in xlsx:
        read(path.read_bytes(), path.name, 3040 if '11_11_2_' in path.name else 3038)
    archives = list(dropbox.glob('20261021_125531_*Downloads.zip'))
    if len(archives) != 1:
        raise ValueError('expected exact 21.10 Downloads archive')
    with zipfile.ZipFile(archives[0]) as archive:
        for name in archive.namelist():
            if name.endswith('.xlsx'):
                read(archive.read(name), archives[0].name + ':' + name, 3040 if '11.11 - 2' in name else 3038)

    # Source snapshot is newer than the current autumn generation's start. Earlier
    # September snapshots with the same human-readable title MUST NOT be used.
    rows = db.query("""SELECT account,nm_id,plan_price,source_file FROM wb_promo_plan_price
        WHERE promo_name='Осенние скидки (автоматические скидки)' AND valid_from='2026-10-08'
          AND source_file LIKE '%%08.10.2026%%' ORDER BY account,nm_id""")
    if len(rows) != 29122:
        raise ValueError(f'autumn source unexpectedly changed: {len(rows)}')
    counts = {}
    for row in rows:
        account = row['account']
        # Diskver's two calendar entries have disjoint product exports and identical
        # name/windows. Preserve that ambiguity explicitly, never invent a member ID.
        ids = [3012] if account == 'wb_acc1' else [3010, 3012]
        offer(account, row['nm_id'], row['plan_price'], ids, row['source_file'])
        counts[account] = counts.get(account, 0) + 1
    sources.append({'name': 'DB autumn 08.10 current generation', 'rows': counts,
                    'sha256': hashlib.sha256(json.dumps(rows, default=str, sort_keys=True).encode()).hexdigest()})
    result = {'version': 1, 'prepared': datetime.now(timezone.utc).isoformat(),
              'notes': ['Diskver Autumn IDs are a same-window group; individual membership is not exposed by WB API',
                        'Native auto-discount blocking explicitly skipped by user'],
              'sources': sources, 'offers': offers}
    result['sha256'] = hashlib.sha256(json.dumps(result, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    atomic_json(output, result)
    print(json.dumps({'cards': {a: len(v) for a, v in offers.items()}, 'sources': len(sources), 'sha256': result['sha256']}))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--calendar', required=True)
    parser.add_argument('--out', required=True)
    parser.add_argument('--dropbox', default='/opt/mp-analytics/dropbox')
    args = parser.parse_args()
    build(Path(args.dropbox), json.loads(Path(args.calendar).read_text()), Path(args.out))


if __name__ == '__main__':
    main()
