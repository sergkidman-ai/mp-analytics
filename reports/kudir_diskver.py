# поток: fin
"""Import Diskver's annual XLSX without executing or repairing its formulas."""
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from functools import lru_cache
import hashlib
import json
from pathlib import Path
import re

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
FILENAME = "20261014_141442_Natalie_Pro_natalya_КДР_2026_1_.xlsx"
VERIFICATION_PATH = BASE_DIR / "reports/data/kudir/manual_verifications/diskver_2026.json"


def source_path():
    archived = BASE_DIR / "reports/data/kudir/originals" / FILENAME
    return archived if archived.is_file() else BASE_DIR / "dropbox" / FILENAME


def _money(value):
    if value is None or value in ("", "-", "—"):
        return Decimal("0.00")
    if isinstance(value, str):
        value = value.replace(" ", "").replace(" ", "").replace(",", ".")
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _category(text):
    text = text.lower()
    if text == "доходы от реализации товаров и услуг":
        return "Продажи — площадка не указана"
    if "агентский отчет" in text or "агентский отчёт" in text:
        return "Агентские отчёты — площадка не указана"
    if "себестоимость вб" in text:
        return "Себестоимость WB"
    if "предоплат" in text or "авансовый платеж" in text:
        return "Авансы и предоплаты — проверить признание"
    if "расходы на ведение бизнеса" in text:
        return "Расходы — состав не расшифрован"
    if "процент" in text and ("уплачены проценты" in text or "выплата процентов" in text):
        return "Банковские проценты"
    if "ндфл" in text:
        return "НДФЛ"
    if "страховой взнос" in text:
        return "Страховые взносы"
    return "Прочие операции — сверить первичку"


def _parse(sheet, formulas, quarter):
    operations = []
    footer = None
    # The fixed workbook uses B–F for Section I; other columns are helper formulas.
    if "Доходы," not in str(sheet.cell(5, 5).value) and not any(
        "учитываемые" in str(sheet.cell(r, 5).value) for r in range(1, 7)
    ):
        raise ValueError("Неизвестная структура КУДиР Дисквера")
    for row in range(7, sheet.max_row + 1):
        label = str(sheet.cell(row, 2).value or "")
        if label.startswith("Итого за"):
            footer = row
            break
        description = str(sheet.cell(row, 4).value or "").strip()
        if not description:
            continue
        if any(sheet.cell(row, col).data_type == "e" for col in (5, 6)):
            raise ValueError("Ошибка Excel в сумме операции КУДиР")
        income, expense = (_money(sheet.cell(row, col).value) for col in (5, 6))
        if not income and not expense:
            continue
        document = str(sheet.cell(row, 3).value or "")
        match = re.search(r"\d{2}\.\d{2}\.\d{4}", document)
        date = datetime.strptime(match[0], "%d.%m.%Y").date().isoformat() if match else None
        number = sheet.cell(row, 2)
        outside = bool(date and (
            not date.startswith("2026-") or
            int(date[5:7]) not in range((quarter - 1) * 3 + 1, quarter * 3 + 1)
        ))
        operations.append({
            "row": row, "number": str(number.value or ""),
            "number_error": number.data_type == "e",
            "document": document, "document_date": date,
            "document_inherited": False, "date_outside_quarter": outside,
            "description": description, "income": str(income), "expense": str(expense),
            "category": _category(description),
        })
    if footer is None:
        raise ValueError("Не найден итог квартала КУДиР Дисквера")
    actual = {k: str(sum((Decimal(o[k]) for o in operations), Decimal("0.00"))) for k in ("income", "expense")}
    totals, delta = {"row": footer}, {}
    errors = {}
    for key, col in (("income", 5), ("expense", 6)):
        cell = sheet.cell(footer, col)
        error = cell.data_type == "e"
        totals[key] = None if error or cell.value is None else str(_money(cell.value))
        errors[key] = str(cell.value) if error else None
        delta[key] = None if totals[key] is None else str(Decimal(actual[key]) - Decimal(totals[key]))
        totals[key + "_formula"] = formulas.cell(footer, col).value
    totals["errors"] = errors
    groups = defaultdict(lambda: {"count": 0, "income": Decimal(0), "expense": Decimal(0)})
    for op in operations:
        g = groups[op["category"]]
        g["count"] += 1
        for key in ("income", "expense"):
            g[key] += Decimal(op[key])
    dates = [o["document_date"] for o in operations if o["document_date"]]
    return {
        "quarter": quarter, "year": 2026, "org": "Дисквер", "org_inn": "7811803918",
        "org_source": "Подтверждено титульным листом исходной книги",
        "operations": operations, "accountant": totals, "calculated": actual, "delta": delta,
        "income_rows_outside_formula": [],
        "outside_document_dates": [o["row"] for o in operations if o["date_outside_quarter"]],
        "groups": [{"category": k, "count": g["count"], "income": str(g["income"]), "expense": str(g["expense"])} for k, g in groups.items()],
        "status": "accountant_import_excel_errors" if any(errors.values()) else "accountant_import",
        "excel_errors": dict(Counter(c.value for cells in sheet for c in cells if c.data_type == "e")),
        "is_partial": quarter == 4,
        "document_date_range": {"from": min(dates) if dates else None, "to": max(dates) if dates else None},
        "sheet": sheet.title, "file_type": "xlsx", "file_scope": "year",
    }


@lru_cache(maxsize=2)
def _load_book(mtime_ns, size, verification_mtime_ns):
    path = source_path()
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    try:
        title = " ".join(str(c.value or "") for cells in values["Титульный лист"] for c in cells)
        if "7811803918" not in title.replace(" ", "") or "2026" not in title:
            raise ValueError("Не подтверждён ИНН/год КУДиР Дисквера")
        checksum = hashlib.sha256(path.read_bytes()).hexdigest()
        verification = json.loads(VERIFICATION_PATH.read_text(encoding="utf-8")) if verification_mtime_ns else {}
        verified_source = (verification.get("org") == "diskver" and verification.get("year") == 2026
                           and verification.get("source_sha256") == checksum)
        result = {}
        for quarter in range(1, 5):
            name = f"Доходы и расходы за {quarter} квартал"
            data = _parse(values[name], formulas[name], quarter)
            data.update(filename=path.name, sha256=checksum)
            checked = verification.get("quarters", {}).get(str(quarter)) if verified_source else None
            if checked == data["calculated"]:
                data["manual_verification"] = {
                    "income": checked["income"], "expense": checked["expense"],
                    "confirmed_by": verification["confirmed_by"],
                    "confirmed_at": verification["confirmed_at"],
                    "source_sha256": checksum,
                }
                data["status"] = "accountant_import_manually_verified"
            result[quarter] = data
        return result
    finally:
        values.close()
        formulas.close()


def load(quarter):
    if quarter not in range(1, 5):
        raise KeyError(quarter)
    stat = source_path().stat()
    verification_mtime = VERIFICATION_PATH.stat().st_mtime_ns if VERIFICATION_PATH.is_file() else 0
    return _load_book(stat.st_mtime_ns, stat.st_size, verification_mtime)[quarter]
